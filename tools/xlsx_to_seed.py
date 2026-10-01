"""Converte Registro Presenze.xlsx in seed.js (dati iniziali dell'app) ed expected.json (valori Excel per i test)."""
import openpyxl, json, re, sys
src = sys.argv[1]
wf = openpyxl.load_workbook(src); wv = openpyxl.load_workbook(src, data_only=True)
hhmm = lambda t: t.strftime('%H:%M') if t else None
years, used, exp = {}, set(), {}
for name in wf.sheetnames:
    if not name.isdigit(): continue
    f, v = wf[name], wv[name]; y = int(name)
    leg = lambda col: f.cell(24, col).value
    cong = []
    for r, col in ((9, 7), (10, 9), (11, 11)):
        lab = (f.cell(r, 37).value or '').replace(' residuo', '')
        a = f.cell(r, 39).value
        m = re.match(r'=(\d+)-', a) if isinstance(a, str) else None
        ini = int(m.group(1)) if m else (a if isinstance(a, (int, float)) else 0)
        cong.append({'code': leg(col), 'label': lab, 'initial': ini})
    pm = int(re.match(r'=(\d+)-', f['AM8'].value).group(1))
    days, ex = {}, []
    for i in range(12):
        b = 8 + 24 * i
        for k, c in enumerate(range(3, 34)):
            d = v.cell(b + 2, c).value
            if d.month != i + 1: continue
            codes = [str(f.cell(r, c).value).strip() for r in (b + 3, b + 4, b + 5) if f.cell(r, c).value not in (None, '')]
            perm = [[hhmm(f.cell(b + a, c).value), hhmm(f.cell(b + a + 1, c).value)] for a in (7, 9) if f.cell(b + a, c).value and f.cell(b + a + 1, c).value]
            ot = [[hhmm(f.cell(b + a, c).value), hhmm(f.cell(b + a + 1, c).value)] for a in (12, 14) if f.cell(b + a, c).value and f.cell(b + a + 1, c).value]
            if codes or perm or ot:
                days[d.strftime('%Y-%m-%d')] = {'codes': codes, 'perm': perm, 'ot': ot, 'note': ''}
                used.update(codes)
        ex.append([v.cell(b + j, 39).value for j in range(0, 10)])
    years[str(y)] = {'archived': False, 'params': {'permMin': pm, 'congedi': cong, 'ferieDisp': None}, 'days': days}
    exp[str(y)] = ex
legend = {'P': ('Permessi', 'Permessi'), 'S': ('Straordinario', 'Straordinario'), 'T': ('Trasferta', 'Trasferta'), 'M': ('Malattia', 'Malattia'),
          'I': ('Infortunio', 'Infortunio'), 'FN': ('Festività Nazionale', 'Festività'), 'FS': ('Festività Soppresse', 'Festività')}
pal = ['#2563eb', '#16a34a', '#d97706', '#dc2626', '#7c3aed', '#0891b2', '#db2777', '#65a30d', '#ea580c', '#475569', '#0d9488', '#9333ea']
allc = list(dict.fromkeys(list(legend) + [c['code'] for y in years.values() for c in y['params']['congedi']] + sorted(used)))
codes = []
for i, c in enumerate(allc):
    d, cat = legend.get(c, ('', 'Congedo' if any(c == x['code'] for y in years.values() for x in y['params']['congedi']) else 'Da definire'))
    codes.append({'code': c, 'desc': d, 'cat': cat, 'color': pal[i % len(pal)], 'val': None, 'active': True})
seed = {'version': 1, 'settings': {'theme': 'auto', 'weekStart': 1, 'lastModified': None, 'lastExport': None}, 'codes': codes, 'years': years}
open('seed.js', 'w').write('window.SEED=' + json.dumps(seed, ensure_ascii=False) + ';')
json.dump(exp, open('/home/claude/expected.json', 'w'), default=str)
print([c['code'] for c in codes], {y: len(v['days']) for y, v in years.items()})
